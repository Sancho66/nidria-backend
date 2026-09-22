"""Import values survive parsing, preview, global resolution and persistence."""

import base64
import io
from pathlib import Path
from typing import Any

import openpyxl
import pytest
import pytest_asyncio
from httpx import AsyncClient
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from shared.models.agent import Agent
from shared.models.client_profile import ClientProfile
from shared.models.custom_field import CustomFieldDefinition
from shared.models.expat_user import ExpatUser
from shared.models.invitation import CaseInvitation
from shared.models.rbac import Role
from src.core.email import outbox
from src.imports.csv_reader import parse_upload
from src.imports.value_resolution import coerce_import_field, country_value
from tests.plugins.agent_plugin import AuthHeaders, MakeAgent

HEADERS = [
    "Prénom du lead",
    "Nom du lead",
    "Origine du contact",
    "Job",
    "Pays de naissance",
    "Pays de résidence",
    "Date et lieu du dernier échange",
    "Etat d'avancement",
    "Nature du lien",
    "Statut de suivi",
]


def leads_file(count: int = 64) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Leads"
    sheet.append(HEADERS)
    for index in range(count):
        sheet.append(
            [
                f"Synthetic{index}",
                "Contact",
                "Rencontre",
                "Consultant",
                "France",
                "Suisse / Paraguay",
                "21 septembre, salon",
                "Premier échange",
                "Connecteur/ Partenaire",
                "relance prevue",
            ]
        )
    sheet.add_table(Table(displayName="Contacts", ref=f"A1:H{count + 1}"))
    sheet["L1"] = "Légende : Profil cible Nidria"
    sheet["M1"] = "Légende : Statut contact Nidria"
    sheet["L66"] = "Legend only, not a contact"
    sheet["Y1000"].fill = PatternFill("solid", fgColor="FFFF00")
    book.create_sheet("Prospects cibles Nidria").append(["Do not import"])
    book.create_sheet("Résumé").append(["Do not import"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_formatted_blank_rows_and_table_boundary() -> None:
    parsed = parse_upload("leads.xlsx", leads_file())
    assert parsed.headers == HEADERS
    assert len(parsed.rows) == 64
    assert parsed.row_indices == list(range(1, 65))
    assert list(parsed.source_rows.values()) == list(range(2, 66))
    assert all(row["Nature du lien"] == "Connecteur/ Partenaire" for row in parsed.rows)
    assert all(row["Statut de suivi"] == "relance prevue" for row in parsed.rows)


def test_provided_workbook_read_only() -> None:
    path = next(Path("samples").glob("Leads_Nidria_Alexia*.xlsx"), None)
    if path is None:
        pytest.skip("Private workbook is not committed; synthetic regression remains mandatory")
    parsed = parse_upload(path.name, path.read_bytes())
    assert parsed.headers == HEADERS
    assert len(parsed.rows) == 64
    assert all(row[HEADERS[0]] and row[HEADERS[1]] for row in parsed.rows)
    assert all(row[HEADERS[8]] and row[HEADERS[9]] for row in parsed.rows)
    assert list(parsed.source_rows.values()) == list(range(2, 66))


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        ("France", "FR"),
        (" fr ", "FR"),
        ("Indonésie", "ID"),
        ("Thailande", "TH"),
        ("Suisse / Paraguay", "Suisse / Paraguay"),
        ("Canada anglais", "Canada anglais"),
    ],
)
def test_country_names_without_guessing_or_truncation(source: str, expected: str) -> None:
    assert country_value(source) == expected


def test_select_labels_and_multiselect_use_canonical_options() -> None:
    definition = CustomFieldDefinition(
        field_type="select", options=["À contacter", "Relance prévue"]
    )
    assert coerce_import_field(definition, "relance PREVUE") == "Relance prévue"
    definition.field_type = "multi_select"
    definition.options = ["Connecteur/ Partenaire", "Influenceur"]
    assert coerce_import_field(definition, "Connecteur/ Partenaire") == ["Connecteur/ Partenaire"]
    assert coerce_import_field(definition, "Influenceur; Connecteur/ Partenaire") == [
        "Influenceur",
        "Connecteur/ Partenaire",
    ]
    definition.options = ["Same", "SAME"]
    with pytest.raises(ValueError):
        coerce_import_field(definition, "same")


@pytest_asyncio.fixture
async def admin(make_agent: MakeAgent, system_roles: dict[str, Role]) -> Agent:
    return await make_agent(role=system_roles["admin"])


@pytest.fixture
async def body(db_session: AsyncSession, admin: Agent) -> dict[str, Any]:
    for key, label, kind, options in [
        ("residence_countries", "Pays de résidence", "text", None),
        ("link_nature", "Nature du lien", "select", ["connector", "influencer"]),
        ("follow_up", "Statut de suivi", "select", ["À contacter", "Relance prévue"]),
    ]:
        db_session.add(
            CustomFieldDefinition(
                agency_id=admin.agency_id,
                key=key,
                label=label,
                field_type=kind,
                scope="person",
                options=options,
            )
        )
    await db_session.commit()
    return {
        "filename": "synthetic-leads.xlsx",
        "file_b64": base64.b64encode(leads_file()).decode(),
        "mapping": dict(
            zip(
                [HEADERS[i] for i in [0, 1, 3, 4, 5, 8, 9]],
                [
                    "first_name",
                    "last_name",
                    "profession",
                    "birth_country",
                    "residence_countries",
                    "link_nature",
                    "follow_up",
                ],
                strict=True,
            )
        ),
        "create_fields": [
            {"column": HEADERS[i], "label": HEADERS[i], "kind": "text"} for i in [2, 6, 7]
        ],
        "require_valid_values": True,
    }


@pytest.mark.usefixtures("rbac_baseline")
async def test_preview_bulk_resolution_and_all_values_stored(
    client: AsyncClient,
    db_session: AsyncSession,
    admin: Agent,
    agent_headers: AuthHeaders,
    body: dict[str, Any],
) -> None:
    headers = agent_headers(admin)
    preview = await client.post(
        "/imports/client-profiles/preview", headers=headers, json={**body, "page_size": 1}
    )
    assert preview.status_code == 200, preview.text
    data = preview.json()
    assert data["total_rows"] == 64
    assert data["summary"]["ignore"] == 0
    row = data["rows"][0]
    assert row["source_row"] == 2
    assert row["source_values"]["link_nature"] == "Connecteur/ Partenaire"
    assert "link_nature" not in row["person"]
    assert row["person"]["birth_country"] == "FR"
    assert row["person"]["residence_countries"] == "Suisse / Paraguay"
    assert row["person"]["follow_up"] == "Relance prévue"
    assert data["value_problems"] == [
        {
            "target": "link_nature",
            "column": "Nature du lien",
            "source_value": "Connecteur/ Partenaire",
            "occurrences": 64,
            "options": ["connector", "influencer"],
        }
    ]
    rejected = await client.post("/imports/client-profiles", headers=headers, json=body)
    assert rejected.status_code == 422
    assert (await db_session.scalar(select(func.count()).select_from(ClientProfile))) == 0
    body["value_mappings"] = [
        {"target": "link_nature", "source_value": "Connecteur/ Partenaire", "value": "connector"}
    ]
    resolved = await client.post(
        "/imports/client-profiles/preview",
        headers=headers,
        json={**body, "page": 2, "page_size": 1},
    )
    assert resolved.status_code == 200, resolved.text
    assert resolved.json()["value_problems"] == []
    assert resolved.json()["rows"][0]["source_row"] == 3
    assert resolved.json()["rows"][0]["person"]["link_nature"] == "connector"
    result = await client.post("/imports/client-profiles", headers=headers, json=body)
    assert result.status_code == 200, result.text
    assert result.json()["created_count"] == 64
    assert [r["source_row"] for r in result.json()["created"]] == list(range(2, 66))
    profiles = (
        await db_session.scalars(
            select(ClientProfile).where(ClientProfile.agency_id == admin.agency_id)
        )
    ).all()
    assert len(profiles) == 64
    for profile in profiles:
        assert profile.profession == "Consultant"
        assert profile.custom_fields == {
            "birth_country": "FR",
            "residence_countries": "Suisse / Paraguay",
            "link_nature": "connector",
            "follow_up": "Relance prévue",
            "origine_du_contact": "Rencontre",
            "date_et_lieu_du_dernier_echange": "21 septembre, salon",
            "etat_d_avancement": "Premier échange",
        }
        assert profile.expat_user_id is None
    assert (await db_session.scalar(select(func.count()).select_from(ExpatUser))) == 0
    assert (await db_session.scalar(select(func.count()).select_from(CaseInvitation))) == 0
    assert outbox == []


@pytest.mark.usefixtures("rbac_baseline")
async def test_blank_rows_differ_from_missing_identity_and_keep_excel_locations(
    client: AsyncClient,
    admin: Agent,
    agent_headers: AuthHeaders,
) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    for row in [
        ["First", "Last", "Job"],
        ["Synthetic", "One", ""],
        [None, "  ", None],
        ["", "", "Consultant"],
        ["Synthetic", "Two", ""],
    ]:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    request = {
        "filename": "rows.xlsx",
        "file_b64": base64.b64encode(buffer.getvalue()).decode(),
        "mapping": {"First": "first_name", "Last": "last_name", "Job": "profession"},
        "corrections": [{"row_index": 4, "target": "first_name", "value": "Corrected"}],
    }
    response = await client.post(
        "/imports/client-profiles/preview", headers=agent_headers(admin), json=request
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["total_rows"] == 3
    assert data["summary"]["ignore_reasons"] == {"missing_identity": 1}
    assert [(row["row_index"], row["source_row"]) for row in data["rows"]] == [
        (1, 2),
        (3, 4),
        (4, 5),
    ]
    assert data["rows"][-1]["person"]["first_name"] == "Corrected"


@pytest.mark.usefixtures("rbac_baseline")
async def test_residence_suggestion_has_no_tax_or_nationality(
    client: AsyncClient,
    admin: Agent,
    agent_headers: AuthHeaders,
) -> None:
    response = await client.post(
        "/imports/client-profiles/suggest-mapping",
        headers=agent_headers(admin),
        json={"headers": ["Pays de résidence"]},
    )
    assert response.status_code == 200, response.text
    assert response.json()["suggestions"] == {}
    assert response.json()["ambiguous"]["Pays de résidence"] == ["residence_address.country"]


@pytest.mark.usefixtures("rbac_baseline")
async def test_value_rules_are_scoped_and_row_corrections_take_precedence(
    client: AsyncClient,
    db_session: AsyncSession,
    admin: Agent,
    make_agent: MakeAgent,
    system_roles: dict[str, Role],
    agent_headers: AuthHeaders,
    body: dict[str, Any],
) -> None:
    rule = {"target": "link_nature", "source_value": "Connecteur/ Partenaire", "value": "connector"}
    resolved = await client.post(
        "/imports/client-profiles/preview",
        headers=agent_headers(admin),
        json={
            **body,
            "value_mappings": [rule],
            "corrections": [{"row_index": 2, "target": "link_nature", "value": "influencer"}],
        },
    )
    assert resolved.status_code == 200, resolved.text
    rows = resolved.json()["rows"]
    assert rows[0]["person"]["link_nature"] == "connector"
    assert rows[1]["person"]["link_nature"] == "influencer"
    assert rows[1]["source_values"]["link_nature"] == "Connecteur/ Partenaire"
    assert rows[2]["person"]["link_nature"] == "connector"
    original = await client.post(
        "/imports/client-profiles/preview",
        headers=agent_headers(admin),
        json=body,
    )
    assert original.json()["value_problems"][0]["occurrences"] == 64
    other = await make_agent(role=system_roles["admin"])
    db_session.add(
        CustomFieldDefinition(
            agency_id=other.agency_id,
            key="link_nature",
            label="Nature",
            field_type="select",
            scope="person",
            options=["different-option"],
        )
    )
    await db_session.commit()
    response = await client.post(
        "/imports/client-profiles/preview",
        headers=agent_headers(other),
        json={
            "csv_text": "First,Last,Nature\nSynthetic,Other,Connecteur/ Partenaire\n",
            "mapping": {"First": "first_name", "Last": "last_name", "Nature": "link_nature"},
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["value_problems"][0]["options"] == ["different-option"]
    assert "link_nature" not in response.json()["rows"][0]["person"]
    invalid = await client.post(
        "/imports/client-profiles/preview",
        headers=agent_headers(admin),
        json={**body, "value_mappings": [{**rule, "target": "unmapped_private_field"}]},
    )
    assert invalid.status_code == 422


@pytest.mark.usefixtures("rbac_baseline")
async def test_compound_country_stays_visible_and_strict_import_writes_nothing(
    client: AsyncClient,
    db_session: AsyncSession,
    admin: Agent,
    agent_headers: AuthHeaders,
) -> None:
    body = {
        "csv_text": "First,Last,Country\nSynthetic,Multiple,Suisse / Paraguay\n",
        "mapping": {
            "First": "first_name",
            "Last": "last_name",
            "Country": "residence_address.country",
        },
        "require_valid_values": True,
    }
    preview = await client.post(
        "/imports/client-profiles/preview", headers=agent_headers(admin), json=body
    )
    assert preview.status_code == 200, preview.text
    row = preview.json()["rows"][0]
    assert row["source_values"]["residence_address.country"] == "Suisse / Paraguay"
    assert row["issues"][0]["source_value"] == "Suisse / Paraguay"
    assert "tax_residence_country" not in row["person"]
    assert preview.json()["value_problems"][0]["occurrences"] == 1
    result = await client.post("/imports/client-profiles", headers=agent_headers(admin), json=body)
    assert result.status_code == 422
    assert result.json()["code"] == "import.unresolved_values"
    assert (await db_session.scalar(select(func.count()).select_from(ClientProfile))) == 0
    assert (await db_session.scalar(select(func.count()).select_from(CustomFieldDefinition))) == 0
